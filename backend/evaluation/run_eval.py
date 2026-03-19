from __future__ import annotations

import argparse
import csv
import json
import os
import re
import statistics
import sys
import unicodedata
import urllib.error
import urllib.request
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

EXPECTED_THEMES = {"manual", "disaster", "remedy"}
QUESTION_FIELDS = {
    "llm_question": "Question",
    "human_question": "Human_question",
}
REQUIRED_COLUMNS = {"Ground_truth", "Evidence", "Source"}
DEFAULT_JUDGE_PROMPT = """Please evaluate these answers based on their accuracy and relevance to the provided passage that based on the Criteria:
1. The Answer is Correct concerning the Reference Answer. Do you agree or disagree? Determine if the given answer accurately matches the reference answer provided. The correctness here means the answer must directly correspond to the reference answer, ensuring factual accuracy.
2. The Answer Includes Relevant, Additional Information from the Context. Do you agree or disagree? Assess whether the answer provides extra details that are not only correct but also relevant and enhance the understanding of the topic as per the information given in the context.
3. The Answer Includes Additional, Irrelevant Information from the Context. Do you agree or disagree? Check if the answer contains extra details that, while related to the context, do not directly pertain to the question asked. This information is not necessary for answering the question and is considered a digression.
4. The Answer Includes Information Not Found in the Context. Do you agree or disagree? Evaluate if the answer includes any information that is not included in the context. This information, even if correct, is extraneous as it goes beyond the provided text and may indicate conjecture or assumption.
5. The Answer is Concise and Free of Redundancy. Do you agree or disagree? Evaluate the efficiency of the language used. Check if the response is direct and avoids repetitive explanations or unnecessary conversational filler. A concise answer should maximize the information-to-word ratio."""


@dataclass
class EvalSettings:
    OPENROUTER_API_KEY: str
    EVAL_CHAT_ENDPOINT: str = "http://localhost:8080/api/v1/chat/"
    EVAL_JUDGE_ENDPOINT: str = "https://openrouter.ai/api/v1/chat/completions"
    EVAL_JUDGE_MODEL: str = "google/gemini-2.5-flash"
    EVAL_JUDGE_TEMPERATURE: float = 0.0


class JudgeDecision(BaseModel):
    correctness: bool = Field(description="True if the answer is correct with respect to the reference.")
    helpfulness: bool = Field(
        description="True if the answer adds relevant details grounded in the context."
    )
    irrelevancy: bool = Field(
        description="True if the answer adds irrelevant details from the context."
    )
    extraneousness: bool = Field(
        description="True if the answer contains information not found in the context."
    )
    conciseness: bool = Field(
        description="True if the answer is concise and avoids redundancy."
    )
    rationale: str = Field(description="A brief explanation for the judgments.")


def load_dotenv_file(env_path: Path) -> None:
    if not env_path.exists():
        return
    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)


def get_eval_settings() -> EvalSettings:
    load_dotenv_file(Path(".env"))
    api_key = normalize_text(os.getenv("OPENROUTER_API_KEY"))
    if not api_key:
        raise ValueError("OPENROUTER_API_KEY is required in environment or .env")

    chat_endpoint = normalize_text(os.getenv("EVAL_CHAT_ENDPOINT")) or "http://localhost:8080/api/v1/chat/"
    judge_endpoint = (
        normalize_text(os.getenv("EVAL_JUDGE_ENDPOINT"))
        or "https://openrouter.ai/api/v1/chat/completions"
    )
    judge_model = normalize_text(os.getenv("EVAL_JUDGE_MODEL")) or "google/gemini-2.5-flash"

    raw_temp = normalize_text(os.getenv("EVAL_JUDGE_TEMPERATURE"))
    judge_temperature = float(raw_temp) if raw_temp else 0.0

    return EvalSettings(
        OPENROUTER_API_KEY=api_key,
        EVAL_CHAT_ENDPOINT=chat_endpoint,
        EVAL_JUDGE_ENDPOINT=judge_endpoint,
        EVAL_JUDGE_MODEL=judge_model,
        EVAL_JUDGE_TEMPERATURE=judge_temperature,
    )


@dataclass
class EvalRow:
    theme: str
    question_type: str
    testcase_id: str
    question: str
    human_question: str
    evaluated_prompt: str
    ground_truth: str
    evidence: str
    source: str
    model_response: str
    retrieved_sources: list[dict[str, str]]
    precision: float
    recall: float
    f1_score: float
    correctness: int
    helpfulness: int
    irrelevancy: int
    extraneousness: int
    conciseness: int
    judge_rationale: str

    def to_dict(self) -> dict[str, Any]:
        payload = {
            "theme": self.theme,
            "question_type": self.question_type,
            "testcase_id": self.testcase_id,
            "Question": self.question,
            "Human_question": self.human_question,
            "evaluated_prompt": self.evaluated_prompt,
            "Ground_truth": self.ground_truth,
            "Evidence": self.evidence,
            "Source": self.source,
            "model_response": self.model_response,
            "precision": round(self.precision, 6),
            "recall": round(self.recall, 6),
            "f1_score": round(self.f1_score, 6),
            "correctness": self.correctness,
            "helpfulness": self.helpfulness,
            "irrelevancy": self.irrelevancy,
            "extraneousness": self.extraneousness,
            "conciseness": self.conciseness,
            "judge_rationale": self.judge_rationale,
            "retrieved_sources": json.dumps(self.retrieved_sources, ensure_ascii=False),
        }
        return payload


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run Phonphai QA evaluation with lexical metrics and LLM-as-a-judge."
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to testcase file (.csv, .json, .jsonl, or .xlsx if openpyxl is installed).",
    )
    parser.add_argument(
        "--output-dir",
        default="evaluation/outputs",
        help="Directory for row-level and summary outputs.",
    )
    parser.add_argument(
        "--chat-endpoint",
        default=None,
        help="Override the chat endpoint. Default comes from EVAL_CHAT_ENDPOINT or .env.",
    )
    parser.add_argument(
        "--judge-model",
        default=None,
        help="Override the judge model. Default comes from EVAL_JUDGE_MODEL or .env.",
    )
    parser.add_argument(
        "--judge-prompt-file",
        default=None,
        help="Optional text file containing the LLM judge instructions.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Only evaluate the first N testcases after loading the dataset.",
    )
    return parser.parse_args()


def load_rows(input_path: Path) -> list[dict[str, Any]]:
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        encodings_to_try = ["utf-8-sig", "utf-8", "utf-16", "cp874", "tis-620", "cp1252", "latin-1"]
        last_error: UnicodeDecodeError | None = None
        for encoding in encodings_to_try:
            try:
                with input_path.open("r", encoding=encoding, newline="") as handle:
                    rows = list(csv.DictReader(handle))
                    if looks_like_mojibake(rows):
                        continue
                    return rows
            except UnicodeDecodeError as exc:
                last_error = exc
        raise RuntimeError(
            f"Unable to decode CSV file '{input_path}' with tried encodings: "
            f"{', '.join(encodings_to_try)}"
        ) from last_error

    if suffix == ".json":
        with input_path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
        if isinstance(payload, list):
            return payload
        raise ValueError("JSON input must be a list of testcase objects.")

    if suffix == ".jsonl":
        rows = []
        with input_path.open("r", encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if line:
                    rows.append(json.loads(line))
        return rows

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "XLSX input requires openpyxl. Install it or convert the testcase file to CSV."
            ) from exc

        workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(cell).strip() if cell is not None else "" for cell in values[0]]
        rows = []
        for row in values[1:]:
            item = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = row[index] if index < len(row) else None
                item[header] = "" if value is None else str(value)
            rows.append(item)
        return rows

    raise ValueError(f"Unsupported input format: {input_path.suffix}")


def validate_rows(rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError("No testcases were loaded from the input file.")

    missing_columns = set()
    for column in QUESTION_FIELDS.values():
        if not any(column in row for row in rows):
            missing_columns.add(column)
    for column in REQUIRED_COLUMNS:
        if not any(column in row for row in rows):
            missing_columns.add(column)
    if missing_columns:
        ordered = ", ".join(sorted(missing_columns))
        raise ValueError(f"Missing required columns: {ordered}")

    unknown_themes = {
        normalize_theme(row.get("theme") or row.get("Theme") or row.get("category"))
        for row in rows
        if row.get("theme") or row.get("Theme") or row.get("category")
    } - EXPECTED_THEMES
    if unknown_themes:
        ordered = ", ".join(sorted(unknown_themes))
        raise ValueError(f"Unsupported theme values found: {ordered}")


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.strip()
    return text


def looks_like_mojibake(rows: list[dict[str, Any]]) -> bool:
    if not rows:
        return False

    sample_values: list[str] = []
    for row in rows[:5]:
        for value in row.values():
            text = normalize_text(value)
            if text:
                sample_values.append(text)
            if len(sample_values) >= 20:
                break
        if len(sample_values) >= 20:
            break

    sample_text = " ".join(sample_values)
    if not sample_text:
        return False

    if contains_thai(sample_text):
        return False

    mojibake_markers = ["Ã", "Ð", "Ñ", "Ò", "ÃÐ", "à", "á", "â", "¤", "º"]
    return any(marker in sample_text for marker in mojibake_markers)


def normalize_theme(value: Any) -> str:
    return normalize_text(value).lower()


def normalize_row(
    row: dict[str, Any],
    index: int,
    forced_theme: str,
) -> dict[str, str]:
    theme = forced_theme
    if theme not in EXPECTED_THEMES:
        raise ValueError(f"Row {index + 1}: theme must be one of {sorted(EXPECTED_THEMES)}.")

    normalized = {
        "theme": theme,
        "Question": normalize_text(row.get("Question")),
        "Human_question": normalize_text(row.get("Human_question")),
        "Ground_truth": normalize_text(row.get("Ground_truth")),
        "Evidence": normalize_text(row.get("Evidence")),
        "Source": normalize_text(row.get("Source")),
        "testcase_id": normalize_text(row.get("testcase_id") or row.get("id") or index + 1),
    }

    if not normalized["Question"] and not normalized["Human_question"]:
        raise ValueError(f"Row {index + 1}: at least one of Question or Human_question is required.")
    if not normalized["Ground_truth"]:
        raise ValueError(f"Row {index + 1}: Ground_truth is required.")

    return normalized


def infer_theme_from_input_path(input_path: Path) -> str:
    input_name = input_path.name.lower()
    for theme in EXPECTED_THEMES:
        if theme in input_name:
            return theme
    return ""


def contains_thai(text: str) -> bool:
    return any("\u0E00" <= char <= "\u0E7F" for char in text)


def tokenize_for_overlap(text: str) -> list[str]:
    normalized = normalize_text(text).lower()
    if not normalized:
        return []

    word_tokens = re.findall(r"[a-z0-9]+", normalized)
    if word_tokens and (not contains_thai(normalized) or len(word_tokens) >= 3):
        return word_tokens

    compact = re.sub(r"\s+", "", normalized)
    if len(compact) < 3:
        return list(compact)
    return [compact[i : i + 3] for i in range(len(compact) - 2)]


def compute_overlap_metrics(prediction: str, reference: str) -> tuple[float, float, float]:
    pred_tokens = tokenize_for_overlap(prediction)
    ref_tokens = tokenize_for_overlap(reference)

    if not pred_tokens and not ref_tokens:
        return 1.0, 1.0, 1.0
    if not pred_tokens or not ref_tokens:
        return 0.0, 0.0, 0.0

    pred_counter = Counter(pred_tokens)
    ref_counter = Counter(ref_tokens)
    overlap = sum((pred_counter & ref_counter).values())

    precision = overlap / sum(pred_counter.values()) if pred_counter else 0.0
    recall = overlap / sum(ref_counter.values()) if ref_counter else 0.0
    if precision + recall == 0:
        f1_score = 0.0
    else:
        f1_score = 2 * precision * recall / (precision + recall)
    return precision, recall, f1_score


def load_judge_prompt(prompt_path: str | None) -> str:
    if not prompt_path:
        return DEFAULT_JUDGE_PROMPT
    return Path(prompt_path).read_text(encoding="utf-8").strip()


def extract_json_object(text: str) -> str:
    text = normalize_text(text)
    if not text:
        raise ValueError("Judge returned an empty response.")

    if text.startswith("```"):
        match = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, flags=re.DOTALL)
        if match:
            return match.group(1)

    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1 or end < start:
        raise ValueError(f"Judge response does not contain JSON: {text}")
    return text[start : end + 1]


def call_openrouter_judge(
    settings: EvalSettings,
    judge_model: str | None,
    prompt: str,
) -> JudgeDecision:
    model_name = judge_model or settings.EVAL_JUDGE_MODEL
    response_schema = {
        "type": "json_schema",
        "json_schema": {
            "name": "judge_decision",
            "strict": True,
            "schema": {
                "type": "object",
                "properties": {
                    "correctness": {"type": "boolean"},
                    "helpfulness": {"type": "boolean"},
                    "irrelevancy": {"type": "boolean"},
                    "extraneousness": {"type": "boolean"},
                    "conciseness": {"type": "boolean"},
                    "rationale": {"type": "string"},
                },
                "required": [
                    "correctness",
                    "helpfulness",
                    "irrelevancy",
                    "extraneousness",
                    "conciseness",
                    "rationale",
                ],
                "additionalProperties": False,
            },
        },
    }
    payload = {
        "model": model_name,
        "temperature": settings.EVAL_JUDGE_TEMPERATURE,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are an evaluation judge. Return only valid JSON that matches the schema."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        "response_format": response_schema,
    }
    request = urllib.request.Request(
        settings.EVAL_JUDGE_ENDPOINT,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {settings.OPENROUTER_API_KEY}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            body = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        error_body = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(
            f"Judge request failed with HTTP {exc.code}: {error_body or exc.reason}"
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Unable to reach judge endpoint '{settings.EVAL_JUDGE_ENDPOINT}'."
        ) from exc

    try:
        content = body["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError(f"Unexpected judge response shape: {body}") from exc

    parsed = json.loads(extract_json_object(content))
    return JudgeDecision.model_validate(parsed)


def call_chat_endpoint(chat_endpoint: str, question: str, testcase_id: str) -> tuple[str, list[dict[str, str]]]:
    payload = {
        "session_id": f"eval-{testcase_id}-{uuid.uuid4().hex[:8]}",
        "message": question,
    }
    request = urllib.request.Request(
        chat_endpoint,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            body = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        error_body = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(
            f"Chat request failed with HTTP {exc.code}: {error_body or exc.reason}"
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Unable to reach chat endpoint '{chat_endpoint}'. Ensure the API gateway is running."
        ) from exc

    answer = normalize_text(body.get("response"))
    sources = body.get("sources") or []
    if not isinstance(sources, list):
        sources = []
    normalized_sources = []
    for item in sources:
        if isinstance(item, dict):
            normalized_sources.append(
                {
                    "title": normalize_text(item.get("title")),
                    "theme": normalize_text(item.get("theme")),
                    "content": normalize_text(item.get("content")),
                }
            )
    return answer, normalized_sources


def judge_answer(
    settings: EvalSettings,
    judge_model: str | None,
    judge_prompt: str,
    question: str,
    ground_truth: str,
    evidence: str,
    source: str,
    answer: str,
) -> JudgeDecision:
    prompt = f"""{judge_prompt}

Question:
{question}

Reference Answer:
{ground_truth}

Context Evidence:
{evidence}

Source Document:
{source}

Candidate Answer:
{answer}
"""
    return call_openrouter_judge(
        settings=settings,
        judge_model=judge_model,
        prompt=prompt,
    )


def mean_or_zero(values: list[float]) -> float:
    if not values:
        return 0.0
    return statistics.fmean(values)


def aggregate_results(results: list[EvalRow]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], list[EvalRow]] = defaultdict(list)
    for row in results:
        grouped[(row.theme, row.question_type)].append(row)

    summary = []
    for (theme, question_type), rows in sorted(grouped.items()):
        summary.append(
            {
                "theme": theme,
                "question_type": question_type,
                "testcase_count": len(rows),
                "precision_mean": round(mean_or_zero([row.precision for row in rows]), 6),
                "recall_mean": round(mean_or_zero([row.recall for row in rows]), 6),
                "f1_score_mean": round(mean_or_zero([row.f1_score for row in rows]), 6),
                "correctness_mean": round(mean_or_zero([row.correctness for row in rows]), 6),
                "helpfulness_mean": round(mean_or_zero([row.helpfulness for row in rows]), 6),
                "irrelevancy_mean": round(mean_or_zero([row.irrelevancy for row in rows]), 6),
                "extraneousness_mean": round(
                    mean_or_zero([row.extraneousness for row in rows]), 6
                ),
                "conciseness_mean": round(mean_or_zero([row.conciseness for row in rows]), 6),
            }
        )

    overall_by_question_type: dict[str, list[EvalRow]] = defaultdict(list)
    for row in results:
        overall_by_question_type[row.question_type].append(row)

    for question_type, rows in sorted(overall_by_question_type.items()):
        summary.append(
            {
                "theme": "overall",
                "question_type": question_type,
                "testcase_count": len(rows),
                "precision_mean": round(mean_or_zero([row.precision for row in rows]), 6),
                "recall_mean": round(mean_or_zero([row.recall for row in rows]), 6),
                "f1_score_mean": round(mean_or_zero([row.f1_score for row in rows]), 6),
                "correctness_mean": round(mean_or_zero([row.correctness for row in rows]), 6),
                "helpfulness_mean": round(mean_or_zero([row.helpfulness for row in rows]), 6),
                "irrelevancy_mean": round(mean_or_zero([row.irrelevancy for row in rows]), 6),
                "extraneousness_mean": round(
                    mean_or_zero([row.extraneousness for row in rows]), 6
                ),
                "conciseness_mean": round(mean_or_zero([row.conciseness for row in rows]), 6),
            }
        )

    return summary


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames = list(rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    args = parse_args()
    settings = get_eval_settings()
    chat_endpoint = args.chat_endpoint or settings.EVAL_CHAT_ENDPOINT
    judge_prompt = load_judge_prompt(args.judge_prompt_file)

    input_path = Path(args.input)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    loaded_rows = load_rows(input_path)
    if args.limit is not None:
        loaded_rows = loaded_rows[: args.limit]
    validate_rows(loaded_rows)

    forced_theme = infer_theme_from_input_path(input_path)
    if not forced_theme:
        raise ValueError(
            "Unable to infer theme from input filename. Include one of: manual, disaster, remedy."
        )
    normalized_rows = [
        normalize_row(row, index, forced_theme=forced_theme)
        for index, row in enumerate(loaded_rows)
    ]
    results: list[EvalRow] = []

    for index, row in enumerate(normalized_rows, start=1):
        print(
            f"[{index}/{len(normalized_rows)}] Evaluating testcase {row['testcase_id']} ({row['theme']})",
            flush=True,
        )
        for question_type, column_name in QUESTION_FIELDS.items():
            prompt = row[column_name]
            if not prompt:
                continue

            answer, retrieved_sources = call_chat_endpoint(
                chat_endpoint=chat_endpoint,
                question=prompt,
                testcase_id=str(row["testcase_id"]),
            )
            precision, recall, f1_score = compute_overlap_metrics(answer, row["Ground_truth"])
            judgment = judge_answer(
                settings=settings,
                judge_model=args.judge_model,
                judge_prompt=judge_prompt,
                question=prompt,
                ground_truth=row["Ground_truth"],
                evidence=row["Evidence"],
                source=row["Source"],
                answer=answer,
            )

            results.append(
                EvalRow(
                    theme=row["theme"],
                    question_type=question_type,
                    testcase_id=str(row["testcase_id"]),
                    question=row["Question"],
                    human_question=row["Human_question"],
                    evaluated_prompt=prompt,
                    ground_truth=row["Ground_truth"],
                    evidence=row["Evidence"],
                    source=row["Source"],
                    model_response=answer,
                    retrieved_sources=retrieved_sources,
                    precision=precision,
                    recall=recall,
                    f1_score=f1_score,
                    correctness=int(judgment.correctness),
                    helpfulness=int(judgment.helpfulness),
                    irrelevancy=int(judgment.irrelevancy),
                    extraneousness=int(judgment.extraneousness),
                    conciseness=int(judgment.conciseness),
                    judge_rationale=judgment.rationale,
                )
            )

    row_output = [row.to_dict() for row in results]
    summary_output = aggregate_results(results)

    rows_csv_path = output_dir / "eval_rows.csv"
    rows_json_path = output_dir / "eval_rows.json"
    summary_csv_path = output_dir / "eval_summary.csv"
    summary_json_path = output_dir / "eval_summary.json"

    write_csv(rows_csv_path, row_output)
    rows_json_path.write_text(
        json.dumps(row_output, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_csv(summary_csv_path, summary_output)
    summary_json_path.write_text(
        json.dumps(summary_output, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"Saved row-level results to {rows_csv_path}")
    print(f"Saved summary results to {summary_csv_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
