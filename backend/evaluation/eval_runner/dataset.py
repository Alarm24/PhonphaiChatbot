from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from .models import EXPECTED_THEMES, QUESTION_FIELDS, REQUIRED_COLUMNS
from .text_utils import contains_thai, normalize_text


def looks_like_mojibake(rows: list[dict[str, Any]]) -> bool:
    if not rows:
        return False

    sample_values: list[str] = []
    for row in rows[:5]:
        for value in row.values():
            text = normalize_text(value)
            if text:
                sample_values.append(text)
            if len(sample_values) >= 20:
                break
        if len(sample_values) >= 20:
            break

    sample_text = " ".join(sample_values)
    if not sample_text or contains_thai(sample_text):
        return False

    mojibake_markers = ["Ã", "Ð", "Ñ", "Ò", "ÃÐ", "à", "á", "â", "¤", "º"]
    return any(marker in sample_text for marker in mojibake_markers)


def load_rows(input_path: Path) -> list[dict[str, Any]]:
    suffix = input_path.suffix.lower()

    if suffix == ".csv":
        encodings_to_try = ["utf-8-sig", "utf-8", "utf-16", "cp874", "tis-620", "cp1252", "latin-1"]
        last_error: UnicodeDecodeError | None = None
        for encoding in encodings_to_try:
            try:
                with input_path.open("r", encoding=encoding, newline="") as handle:
                    rows = list(csv.DictReader(handle))
                    if looks_like_mojibake(rows):
                        continue
                    return rows
            except UnicodeDecodeError as exc:
                last_error = exc
        raise RuntimeError(
            f"Unable to decode CSV file '{input_path}' with tried encodings: "
            f"{', '.join(encodings_to_try)}"
        ) from last_error

    if suffix == ".json":
        with input_path.open("r", encoding="utf-8") as handle:
            payload = json.load(handle)
        if isinstance(payload, list):
            return payload
        raise ValueError("JSON input must be a list of testcase objects.")

    if suffix == ".jsonl":
        rows = []
        with input_path.open("r", encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if line:
                    rows.append(json.loads(line))
        return rows

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "XLSX input requires openpyxl. Install it or convert the testcase file to CSV."
            ) from exc

        workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []

        headers = [str(cell).strip() if cell is not None else "" for cell in values[0]]
        rows = []
        for row in values[1:]:
            item = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                value = row[index] if index < len(row) else None
                item[header] = "" if value is None else str(value)
            rows.append(item)
        return rows

    raise ValueError(f"Unsupported input format: {input_path.suffix}")


def validate_rows(rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError("No testcases were loaded from the input file.")

    missing_columns = set()
    for column in QUESTION_FIELDS.values():
        if not any(column in row for row in rows):
            missing_columns.add(column)
    for column in REQUIRED_COLUMNS:
        if not any(column in row for row in rows):
            missing_columns.add(column)
    if missing_columns:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing_columns))}")


def infer_theme_from_input_path(input_path: Path) -> str:
    input_name = input_path.name.lower()
    for theme in EXPECTED_THEMES:
        if theme in input_name:
            return theme
    return ""


def normalize_row(row: dict[str, Any], index: int, forced_theme: str) -> dict[str, str]:
    if forced_theme not in EXPECTED_THEMES:
        raise ValueError(f"Row {index + 1}: theme must be one of {sorted(EXPECTED_THEMES)}.")

    normalized = {
        "theme": forced_theme,
        "Question": normalize_text(row.get("Question")),
        "Human_question": normalize_text(row.get("Human_question")),
        "Ground_truth": normalize_text(row.get("Ground_truth")),
        "Evidence": normalize_text(row.get("Evidence")),
        "Source": normalize_text(row.get("Source")),
        "testcase_id": normalize_text(row.get("testcase_id") or row.get("id") or index + 1),
    }

    if not normalized["Question"] and not normalized["Human_question"]:
        raise ValueError(
            f"Row {index + 1}: at least one of Question or Human_question is required."
        )
    if not normalized["Ground_truth"]:
        raise ValueError(f"Row {index + 1}: Ground_truth is required.")

    return normalized
